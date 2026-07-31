"""把 20260714_Overview.pptx 的五張表格，逐格重建成 Excel（含合併儲存格、雙層表頭）。

版面與內容與 PPT 完全一致（PPT 原順序：CDR → CASI → MMSE → 1:1 配對 → 全世代人口學）。
數值皆已與真實資料核對一致（見 cognitive_strata_table.py 與 dataset_figure.py 的即時重算）；
本檔為「照 PPT 版面」的靜態重建，故數字以字串常數保存，含 PPT 之原始格式（逗號、FeMale 拼法）。

用法：python scripts/paper/overview_tables_xlsx.py [-o out.xlsx]
"""
import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # scripts/
from _paths import PROJECT_ROOT  # noqa: F401

DEFAULT_OUTPUT = PROJECT_ROOT / "paper" / "draft" / "v5" / "overview_tables.xlsx"

_side = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_side, right=_side, top=_side, bottom=_side)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")

# ── 五張表：grid（"" 代表合併延續格）、merges（表內 1-indexed）、表頭列數、欄寬 ──
CDR = dict(
    sheet="CDR", title="表：Global CDR 分佈（subject-level，人數）",
    n_header=2, widths=[10, 9, 11, 7, 7, 7, 7, 7, 8],
    grid=[
        ["組別", "n（人）", "拍攝（人次）", "CDR", "", "", "", "", "缺測"],
        ["", "", "", "0", "0.5", "1", "2", "3", ""],
        ["P", "1061", "1061", "16", "272", "549", "145", "8", "71"],
        ["SCD", "458", "791", "42", "94", "14", "1", "0", "307"],
        ["ACS", "91", "218", "0", "0", "0", "0", "0", "91"],
        ["HC", "549", "1009", "42", "94", "14", "1", "0", "398"],
    ],
    merges=[(1, 1, 2, 1), (1, 2, 2, 2), (1, 3, 2, 3), (1, 4, 1, 8), (1, 9, 2, 9)],
)
CASI = dict(
    sheet="CASI", title="表：CASI 分層（subject-level，人數；80／50 為暫定切點）",
    n_header=2, widths=[10, 9, 11, 8, 9, 8, 8],
    grid=[
        ["組別", "n（人）", "拍攝（人次）", "CASI", "", "", "缺測"],
        ["", "", "", "≥80", "50–79", "≤49", ""],
        ["P", "1061", "1061", "104", "504", "384", "69"],
        ["SCD", "458", "791", "284", "93", "8", "73"],
        ["ACS", "91", "218", "85", "0", "0", "6"],
        ["HC", "549", "1009", "369", "93", "8", "79"],
    ],
    merges=[(1, 1, 2, 1), (1, 2, 2, 2), (1, 3, 2, 3), (1, 4, 1, 6), (1, 7, 2, 7)],
)
MMSE = dict(
    sheet="MMSE", title="表：MMSE 分層（subject-level，人數）",
    n_header=2, widths=[10, 9, 11, 8, 9, 8, 8],
    grid=[
        ["組別", "n（人）", "拍攝（人次）", "MMSE", "", "", "缺測"],
        ["", "", "", "≥26", "18–25", "≤17", ""],
        ["P", "1061", "1061", "45", "380", "567", "69"],
        ["SCD", "458", "791", "193", "170", "22", "73"],
        ["ACS", "91", "218", "78", "7", "0", "6"],
        ["HC", "549", "1009", "271", "177", "22", "79"],
    ],
    merges=[(1, 1, 2, 1), (1, 2, 2, 2), (1, 3, 2, 3), (1, 4, 1, 6), (1, 7, 2, 7)],
)
MATCHED = dict(
    sheet="1by1_matched", title="表：1 by 1 年齡配對後之兩臂（拍攝層級）",
    n_header=2,
    widths=[12, 7, 11, 11, 16, 16, 9, 9, 11, 11, 16, 16, 9, 9],
    grid=[
        ["配對集", "Pairs", "患者側 (P)", "", "", "", "", "",
         "對照側 (Control)", "", "", "", "", ""],
        ["", "", "人 (Subjects)", "人次 (Visits)", "Age(人次) (mean ± SD)",
         "BMI(人次) (mean ± SD)", "Male (人次)", "FeMale (人次)",
         "人 (Subjects)", "人次 (Visits)", "Age(人次) (mean ± SD)",
         "BMI(人次) (mean ± SD)", "Male (人次)", "FeMale (人次)"],
        ["SCD-matched", "519", "519", "519", "77.5 ± 5.4", "22.8 ± 3.6", "187", "332",
         "290", "519", "77.3 ± 5.5", "23.9 ± 3.9", "202", "317"],
        ["ACS-matched", "77", "77", "77", "66.4 ± 4.3", "23.0 ± 3.8", "33", "44",
         "31", "77", "66.2 ± 4.3", "23.2 ± 2.9", "16", "61"],
        ["HC-matched", "596", "596", "596", "76.1 ± 6.5", "22.8 ± 3.6", "220", "376",
         "321", "596", "75.9 ± 6.5", "23.8 ± 3.8", "218", "378"],
    ],
    merges=[(1, 1, 2, 1), (1, 2, 2, 2), (1, 3, 1, 8), (1, 9, 1, 14)],
)
DEMO = dict(
    sheet="Demographics", title="表：全世代人口學（subject-level；Age／BMI 為每人首訪值）",
    n_header=1,
    widths=[8, 7, 9, 8, 9, 8, 14, 12, 14, 12, 9, 7, 9, 7],
    grid=[
        ["Group", "", "Subjects", "", "Visits", "", "Age(人)(mean±SD)", "",
         "BMI(人)(mean±SD)", "", "Male(人)", "", "Female(人)", ""],
        ["P", "", "1,061", "", "1,061", "", "79.8 ± 7.0", "", "22.9 ± 3.7", "",
         "366", "", "695", ""],
        ["HC", "SCD", "549", "458", "1009", "791", "70.0 ± 9.1", "72.4 ± 7.6",
         "23.7 ± 3.8", "23.7 ± 3.8", "183", "165", "366", "293"],
        ["", "ACS", "", "91", "", "218", "", "57.9 ± 6.1", "", "23.5 ± 3.8",
         "", "18", "", "73"],
        ["Total", "", "1,610", "", "2,070", "", "76.4 ± 9.1", "", "23.2 ± 3.8", "",
         "549", "", "1,061", ""],
    ],
    merges=[
        # header：每個指標橫跨兩子欄
        (1, 1, 1, 2), (1, 3, 1, 4), (1, 5, 1, 6), (1, 7, 1, 8),
        (1, 9, 1, 10), (1, 11, 1, 12), (1, 13, 1, 14),
        # P 列：無子群，橫跨兩子欄
        (2, 1, 2, 2), (2, 3, 2, 4), (2, 5, 2, 6), (2, 7, 2, 8),
        (2, 9, 2, 10), (2, 11, 2, 12), (2, 13, 2, 14),
        # HC 直向跨 SCD/ACS 兩列
        (3, 1, 4, 1),
        # Total 列：橫跨兩子欄
        (5, 1, 5, 2), (5, 3, 5, 4), (5, 5, 5, 6), (5, 7, 5, 8),
        (5, 9, 5, 10), (5, 11, 5, 12), (5, 13, 5, 14),
    ],
)

TABLES = [CDR, CASI, MMSE, MATCHED, DEMO]   # PPT 原順序


def write_table(ws, spec):
    grid, merges, n_header = spec["grid"], spec["merges"], spec["n_header"]
    ncols = len(grid[0])
    # 標題列
    ws.cell(1, 1, spec["title"]).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    r0 = 3  # 空一列後開始畫表
    for i, row in enumerate(grid):
        ws.row_dimensions[r0 + i].height = 30 if i < n_header else 20
        for j, val in enumerate(row):
            c = ws.cell(r0 + i, j + 1, val if val != "" else None)
            c.border = BORDER
            c.alignment = CENTER
            if i < n_header:
                c.font = BOLD
                c.fill = HDR_FILL
    for (a, b, cc, d) in merges:
        ws.merge_cells(start_row=r0 + a - 1, start_column=b,
                       end_row=r0 + cc - 1, end_column=d)
    for idx, w in enumerate(spec["widths"]):
        ws.column_dimensions[get_column_letter(idx + 1)].width = w


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT)
    args = ap.parse_args()

    wb = Workbook()
    wb.remove(wb.active)
    for spec in TABLES:
        write_table(wb.create_sheet(spec["sheet"]), spec)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print("saved ->", args.output, "| sheets:", [s["sheet"] for s in TABLES])


if __name__ == "__main__":
    main()
